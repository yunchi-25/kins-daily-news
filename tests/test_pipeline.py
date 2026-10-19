"""mock 네이버 API + mock 원문 확인으로 전체 파이프라인과 내보내기를 검증한다.

여기의 기사 데이터는 테스트용 fixture 이며 실제 기사가 아니다.
"""

from __future__ import annotations

import io
from datetime import datetime, timedelta

import pytest
from openpyxl import load_workbook

from src import pipeline
from src.config import NEWS_KEYWORDS, TOPICS, TOPIC_DOMESTIC, TOPIC_ETC, TOPIC_KINS, TOPIC_NSSC, Settings
from src.exporter import all_articles_df, keyword_df, to_csv_bytes, to_excel_bytes, to_text_report, topic_df
from src.models import SOURCE_NAVER, TEXT_OK, RawSearchItem
from src.naver_news import KeywordSearchResult, NaverApiError, NaverAuthError, NaverNewsClient
from src.config import KST
from src.views import Filters, keyword_view, topic_view

START = datetime(2026, 10, 16, 11, 0, tzinfo=KST)
END = datetime(2026, 10, 19, 8, 13, 5, tzinfo=KST)

GORI = "원자력안전위원회는 회의를 열고 고리 3호기 계속운전 허가안을 의결했다. 한국수력원자력은 2034년까지 운전한다. " * 4
GORI2 = "고리 3호기가 10년 더 운전한다. 원자력안전위원회는 고리 3호기 계속운전 허가안을 의결했다고 밝혔다. " * 4
KINS = "한국원자력안전기술원(KINS)은 방사선 안전 연구 결과를 발표했다. KINS는 ... " * 4
HANBIT = "한빛원전 1호기에서 냉각재 누설이 발생해 원자로가 수동 정지됐다. 한국수력원자력은 원인을 조사 중이다. " * 4
PUBLIC = "기획재정부는 공공기관 통폐합 방안을 발표했다. 공공기관운영위원회 심의를 거친다. 공공기관 기능 조정. " * 4
STOCK = "원전 관련주가 일제히 상승했다. " * 4

# (keyword, title, original, naver, pubDate, body)
FIXTURE = [
    ("원안위", "원안위, 고리 3호기 계속운전 승인", "https://a.example.com/1", "https://n.news.naver.com/mnews/article/001/0000000001", END - timedelta(hours=10), GORI),
    ("원전", "원안위, 고리 3호기 계속운전 승인", "https://a.example.com/1", "https://n.news.naver.com/mnews/article/001/0000000001", END - timedelta(hours=10), GORI),
    ("원자력안전위원회", "고리3호기 수명 연장…원안위 계속운전 승인", "https://b.example.com/2", None, END - timedelta(hours=9), GORI2),
    ("KINS", "KINS, 방사선 안전 연구 결과 발표", "https://c.example.com/3", None, END - timedelta(hours=5), KINS),
    ("한국원자력안전기술원", "KINS, 방사선 안전 연구 결과 발표", "https://c.example.com/3", None, END - timedelta(hours=5), KINS),
    ("원전", "한빛 1호기 냉각재 누설로 수동 정지", "https://d.example.com/4", None, START, HANBIT),  # 시작 시각 정확히 일치
    ("공공기관", "기재부, 공공기관 통폐합 추진", "https://e.example.com/5", None, END, PUBLIC),  # 종료 시각 정확히 일치
    ("원전", "[특징주] 원전 관련주 급등", "https://f.example.com/6", None, END - timedelta(hours=1), STOCK),
    ("원전", "원전 지난주 기사", "https://g.example.com/7", None, START - timedelta(seconds=1), HANBIT),  # 기간 밖
    ("최원호", "최원호 감독, 시즌 마무리 소감", "https://h.example.com/8", None, END - timedelta(hours=2), "프로야구 최원호 감독이 시즌을 마무리했다. 원전 " * 3),
]


class FakeNaver(NaverNewsClient):
    def __init__(self, fail_keyword: str | None = None, auth_fail: bool = False):
        self.calls: list[str] = []
        self.fail_keyword = fail_keyword
        self.auth_fail = auth_fail

    def search_keyword(self, keyword, start_dt, end_dt, max_results=1000):
        self.calls.append(keyword)
        if self.auth_fail:
            raise NaverAuthError("인증 실패")
        if keyword == self.fail_keyword:
            raise NaverApiError("timeout")
        res = KeywordSearchResult(keyword=keyword)
        for kw, title, orig, naver, pub, _ in FIXTURE:
            if kw == keyword:
                res.items.append(RawSearchItem(kw, title, orig, naver, title, pub))
        res.total_fetched = len(res.items)
        return res


BODIES = {orig: body for _, _, orig, _, _, body in FIXTURE}


def fake_enrich(article, session, now):
    article.article_text = BODIES.get(article.original_url)
    article.text_status = TEXT_OK
    article.published_at = article.api_pub_date
    article.published_at_source = SOURCE_NAVER
    article.publisher = "테스트일보"
    return article


def _settings() -> Settings:
    return Settings(naver_client_id="x", naver_client_secret="y")


@pytest.fixture()
def run():
    progress = []
    r = pipeline.run_search(
        START, END, _settings(), naver_client=FakeNaver(), enrich_fn=fake_enrich,
        progress=lambda s, f, m: progress.append((s, f, m)),
    )
    r.stats["_progress"] = progress
    return r


def test_all_28_keywords_searched_in_order():
    fake = FakeNaver()
    pipeline.run_search(START, END, _settings(), naver_client=fake, enrich_fn=fake_enrich)
    assert fake.calls == NEWS_KEYWORDS


def test_pipeline_counts_and_filters(run):
    s = run.stats
    assert s["api_results"] == len(FIXTURE)
    assert s["collected"] == len(FIXTURE) - 1  # 기간 밖 1건 제외
    titles = {a.title for a in run.included}
    assert "원안위, 고리 3호기 계속운전 승인" in titles
    assert "한빛 1호기 냉각재 누설로 수동 정지" in titles  # 시작 시각 경계 포함
    assert "기재부, 공공기관 통폐합 추진" in titles  # 종료 시각 경계 포함
    assert "[특징주] 원전 관련주 급등" not in titles
    assert "원전 지난주 기사" not in titles
    assert "최원호 감독, 시즌 마무리 소감" not in titles  # 동명이인


def test_exact_duplicates_merged_with_keywords(run):
    gori = [a for a in run.articles if a.title == "원안위, 고리 3호기 계속운전 승인"]
    assert len(gori) == 1
    assert gori[0].matched_keywords == ["원안위", "원전"]


def test_topic_view_one_topic_and_representatives(run):
    view = topic_view(run)
    assert list(view.keys()) == TOPICS
    shown = [a.article_id for items in view.values() for a in items]
    assert len(shown) == len(set(shown))  # 한 기사는 한 주제에만
    # 같은 사건 유사 보도 2건 중 대표(중요도 최고) 1건만 표시, 나머지는 대표에 연결
    gori_titles = {"원안위, 고리 3호기 계속운전 승인", "고리3호기 수명 연장…원안위 계속운전 승인"}
    assert len(view[TOPIC_NSSC]) == 1 and view[TOPIC_NSSC][0].title in gori_titles
    rep = view[TOPIC_NSSC][0]
    assert len(rep.similar_article_ids) == 1
    assert rep.relevance_score >= run.by_id()[rep.similar_article_ids[0]].relevance_score
    assert [a.title for a in view[TOPIC_KINS]] == ["KINS, 방사선 안전 연구 결과 발표"]
    assert [a.title for a in view[TOPIC_DOMESTIC]] == ["한빛 1호기 냉각재 누설로 수동 정지"]
    assert [a.title for a in view[TOPIC_ETC]] == ["기재부, 공공기관 통폐합 추진"]
    assert run.stats["final"] == 4
    assert run.stats["event_merged"] == 1


def test_keyword_view_shows_all_keywords_with_duplicates(run):
    view = keyword_view(run, collapse_similar=False)
    assert list(view.keys()) == NEWS_KEYWORDS  # 0건 키워드 포함 28개 모두
    assert len(view["라돈"]) == 0
    # 같은 기사가 여러 키워드에 표시됨
    assert any(a.title == "원안위, 고리 3호기 계속운전 승인" for a in view["원안위"])
    assert any(a.title == "원안위, 고리 3호기 계속운전 승인" for a in view["원전"])
    assert any(a.title == "KINS, 방사선 안전 연구 결과 발표" for a in view["KINS"])
    assert any(a.title == "KINS, 방사선 안전 연구 결과 발표" for a in view["한국원자력안전기술원"])
    # 비대표 유사 보도도 실제 검색된 키워드에서는 표시
    assert [a.title for a in view["원자력안전위원회"]] == ["고리3호기 수명 연장…원안위 계속운전 승인"]
    assert view["최원호"] == []


def test_filters_do_not_call_api(run):
    v = topic_view(run, Filters(topics=[TOPIC_KINS]))
    assert sum(len(x) for x in v.values()) == 1
    v2 = keyword_view(run, Filters(keywords=["원전"], query="한빛"))
    assert list(v2.keys()) == ["원전"] and len(v2["원전"]) == 1


def test_progress_reported(run):
    msgs = [m for _, _, m in run.stats["_progress"]]
    assert msgs[0] == "검색 준비 중..."
    assert "[ 1 / 28 ] 한국원자력안전기술원 검색 중" in msgs
    assert "[ 28 / 28 ] 석광훈 검색 중" in msgs
    fracs = [f for _, f, _ in run.stats["_progress"]]
    assert fracs == sorted(fracs) and fracs[-1] == 1.0


def test_keyword_failure_does_not_stop_search():
    r = pipeline.run_search(START, END, _settings(), naver_client=FakeNaver(fail_keyword="원전"), enrich_fn=fake_enrich)
    assert any("'원전' 검색 실패" in w for w in r.warnings)
    assert r.stats["final"] >= 1


def test_auth_failure_aborts():
    with pytest.raises(pipeline.SearchAborted):
        pipeline.run_search(START, END, _settings(), naver_client=FakeNaver(auth_fail=True), enrich_fn=fake_enrich)


def test_missing_credentials_message():
    with pytest.raises(pipeline.SearchAborted) as exc:
        pipeline.run_search(START, END, Settings())
    assert "NAVER_CLIENT_ID" in str(exc.value) and "NAVER_CLIENT_SECRET" in str(exc.value)


def test_enrich_exception_does_not_stop_pipeline():
    def broken(article, session, now):
        raise RuntimeError("parse error")

    r = pipeline.run_search(START, END, _settings(), naver_client=FakeNaver(), enrich_fn=broken)
    assert r.stats["unique"] > 0  # 예외가 전파되지 않음


def test_llm_failure_falls_back_to_rules(run):
    from src.llm import LLMError

    class BrokenLLM:
        def judge_articles(self, batch):
            raise LLMError("timeout")

    s = _settings()
    s.openai_api_key, s.openai_model = "k", "m"
    r = pipeline.run_search(START, END, s, naver_client=FakeNaver(), enrich_fn=fake_enrich, llm_client=BrokenLLM())
    assert not r.llm_used
    assert r.stats["final"] == run.stats["final"]
    assert any("LLM" in w for w in r.warnings)


def test_llm_judgement_applied():
    from src.llm import LLMJudgement

    class FakeLLM:
        def judge_articles(self, batch):
            out = {}
            for a in batch:
                include = "공공기관" not in a.title
                out[a.article_id] = LLMJudgement(include, None, 90 if include else 10, None, None if include else "테스트 제외", "test")
            return out

        def judge_same_events(self, pairs):
            return {}

    s = _settings()
    s.openai_api_key, s.openai_model = "k", "m"
    r = pipeline.run_search(START, END, s, naver_client=FakeNaver(), enrich_fn=fake_enrich, llm_client=FakeLLM())
    assert r.llm_used
    assert "기재부, 공공기관 통폐합 추진" not in {a.title for a in r.included}
    excluded = {a.title: a.exclusion_reason for a in r.excluded}
    assert excluded["기재부, 공공기관 통폐합 추진"].startswith("LLM 판단")


def test_zero_results():
    class Empty(FakeNaver):
        def search_keyword(self, keyword, start_dt, end_dt, max_results=1000):
            return KeywordSearchResult(keyword=keyword)

    r = pipeline.run_search(START, END, _settings(), naver_client=Empty(), enrich_fn=fake_enrich)
    assert r.stats["final"] == 0
    assert all(len(v) == 0 for v in topic_view(r).values())
    report = to_text_report(r)
    assert report.count("확인된 기사 없음 (0건)") == 6
    to_excel_bytes(r)  # 0건이어도 Excel 생성 가능


# ---------------------------------------------------------------------------
# 내보내기
# ---------------------------------------------------------------------------
def test_excel_sheets_and_content(run):
    data = to_excel_bytes(run)
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["전체기사", "주제별결과", "키워드별결과"]
    header = [c.value for c in wb["전체기사"][1]]
    for col in ["제목", "언론사", "게시시각", "주제", "검색키워드", "원문URL", "네이버URL"]:
        assert col in header
    kw_rows = list(wb["키워드별결과"].iter_rows(min_row=2, values_only=True))
    kws = {r[1] for r in kw_rows}
    assert kws == set(NEWS_KEYWORDS)  # 0건 키워드도 행으로 표시
    gori_rows = [r for r in kw_rows if r[2] == "원안위, 고리 3호기 계속운전 승인"]
    assert {r[1] for r in gori_rows} == {"원안위", "원전"}  # 키워드별로 여러 행


def test_csv_has_bom_and_titles(run):
    data = to_csv_bytes(all_articles_df(run))
    assert data.startswith(b"\xef\xbb\xbf")
    text = data.decode("utf-8-sig")
    assert "원안위, 고리 3호기 계속운전 승인" in text
    assert len(topic_df(run)) >= 6
    assert len(keyword_df(run)) >= 28


def test_text_report_keeps_original_titles(run):
    report = to_text_report(run)
    assert "■ 원자력안전위원회 (1건)" in report
    for a in run.representatives:
        assert a.title in report and a.link in report
    assert "■ 정부&정치권 (0건)\n확인된 기사 없음 (0건)" in report
