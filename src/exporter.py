"""CSV / Excel / 텍스트 보고서 내보내기."""

from __future__ import annotations

import io

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import NEWS_KEYWORDS, TOPICS
from .models import NewsArticle, SearchRun
from .ranking import sort_articles
from .utils import format_kst
from .views import similar_members, topic_view

NO_ARTICLE_TEXT = "확인된 기사 없음 (0건)"


class ExportError(Exception):
    pass


def _base_row(a: NewsArticle) -> dict:
    return {
        "제목": a.title,
        "언론사": a.display_publisher,
        "게시시각": format_kst(a.published_at),
        "주제": a.topic or "",
        "검색키워드": ", ".join(a.matched_keywords),
        "원문URL": a.original_url or "",
        "네이버URL": a.naver_url or "",
    }


def all_articles_df(run: SearchRun) -> pd.DataFrame:
    order = {t: i for i, t in enumerate(TOPICS)}
    items = sorted(run.included, key=lambda a: (order.get(a.topic, 99), not a.is_representative, -(a.relevance_score or 0)))
    rows = []
    for a in items:
        row = _base_row(a)
        row.update(
            {
                "유사보도그룹": a.duplicate_group,
                "대표기사": "Y" if a.is_representative else "",
                "게시시각출처": a.published_at_source or "",
                "본문상태": a.text_status,
            }
        )
        rows.append(row)
    columns = ["제목", "언론사", "게시시각", "주제", "검색키워드", "원문URL", "네이버URL", "유사보도그룹", "대표기사", "게시시각출처", "본문상태"]
    return pd.DataFrame(rows, columns=columns)


def topic_df(run: SearchRun) -> pd.DataFrame:
    groups = similar_members(run)
    rows = []
    for topic, items in topic_view(run).items():
        if not items:
            rows.append({"주제": topic, "순번": "", "제목": NO_ARTICLE_TEXT})
            continue
        for i, a in enumerate(items, start=1):
            row = {"주제": topic, "순번": i}
            row.update(_base_row(a))
            row["유사보도수"] = max(0, len(groups.get(a.duplicate_group, [a])) - 1)
            rows.append(row)
    columns = ["주제", "순번", "제목", "언론사", "게시시각", "검색키워드", "원문URL", "네이버URL", "유사보도수"]
    return pd.DataFrame(rows, columns=columns)


def keyword_df(run: SearchRun) -> pd.DataFrame:
    """키워드별 결과: 한 기사가 여러 키워드에 해당하면 여러 행."""
    rows = []
    for idx, kw in enumerate(NEWS_KEYWORDS, start=1):
        items = sort_articles([a for a in run.included if kw in a.matched_keywords])
        if not items:
            rows.append({"키워드순번": idx, "키워드": kw, "제목": NO_ARTICLE_TEXT})
            continue
        for a in items:
            row = {"키워드순번": idx, "키워드": kw}
            row.update(_base_row(a))
            row["대표기사"] = "Y" if a.is_representative else ""
            row["유사보도그룹"] = a.duplicate_group
            rows.append(row)
    columns = ["키워드순번", "키워드", "제목", "언론사", "게시시각", "주제", "검색키워드", "원문URL", "네이버URL", "대표기사", "유사보도그룹"]
    return pd.DataFrame(rows, columns=columns)


def to_csv_bytes(df: pd.DataFrame) -> bytes:
    # Excel 에서 한글이 깨지지 않도록 BOM 포함 UTF-8
    return df.to_csv(index=False).encode("utf-8-sig")


_HEADER_FILL = PatternFill("solid", fgColor="0B2E59")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_LINK_FONT = Font(color="1F4E8C", underline="single")
_WIDTHS = {
    "제목": 70, "언론사": 16, "게시시각": 17, "주제": 22, "검색키워드": 28, "원문URL": 40, "네이버URL": 40,
    "키워드": 18, "키워드순번": 9, "순번": 6, "유사보도그룹": 11, "대표기사": 8, "유사보도수": 9,
    "게시시각출처": 11, "본문상태": 12,
}


def to_excel_bytes(run: SearchRun) -> bytes:
    try:
        sheets = {
            "전체기사": all_articles_df(run),
            "주제별결과": topic_df(run),
            "키워드별결과": keyword_df(run),
        }
        buf = io.BytesIO()
        with pd.ExcelWriter(buf, engine="openpyxl") as writer:
            for name, df in sheets.items():
                df.to_excel(writer, sheet_name=name, index=False)
                ws = writer.sheets[name]
                for col_idx, col_name in enumerate(df.columns, start=1):
                    cell = ws.cell(row=1, column=col_idx)
                    cell.fill = _HEADER_FILL
                    cell.font = _HEADER_FONT
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    ws.column_dimensions[get_column_letter(col_idx)].width = _WIDTHS.get(col_name, 14)
                    if col_name in ("원문URL", "네이버URL"):
                        for row_idx in range(2, ws.max_row + 1):
                            c = ws.cell(row=row_idx, column=col_idx)
                            if isinstance(c.value, str) and c.value.startswith("http"):
                                c.hyperlink = c.value
                                c.font = _LINK_FONT
                ws.freeze_panes = "A2"
                if ws.max_row > 1:
                    ws.auto_filter.ref = ws.dimensions
        return buf.getvalue()
    except Exception as exc:
        raise ExportError(f"Excel 파일 생성 중 오류가 발생했습니다: {exc}") from exc


def to_text_report(run: SearchRun) -> str:
    """주제별 결과를 업무 보고용 텍스트로 만든다 (제목·언론사·시각·URL 원문 그대로)."""
    lines = [
        "KINS 일일 언론동향",
        f"검색 범위: {format_kst(run.start)} ~ {format_kst(run.end)} (KST)",
        "",
    ]
    for topic, items in topic_view(run).items():
        lines.append(f"■ {topic} ({len(items)}건)")
        if not items:
            lines.append(NO_ARTICLE_TEXT)
        for i, a in enumerate(items, start=1):
            lines.append(f"{i}. {a.title} ({a.display_publisher}, {format_kst(a.published_at, '%m-%d %H:%M')})")
            if a.link:
                lines.append(f"   {a.link}")
        lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def file_stamp(run: SearchRun) -> str:
    return f"{run.start.strftime('%Y%m%d%H%M')}-{run.end.strftime('%Y%m%d%H%M')}"
